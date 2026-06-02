from __future__ import annotations

import json
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent / "kb_import_selftest"
PDF_DIR = ROOT / "pdf"
XLSX_DIR = ROOT / "xlsx"
ZIP_PATH = ROOT / "kb_import_selftest.zip"
MANIFEST_PATH = ROOT / "manifest.json"


@dataclass
class Case:
    prefix: str
    part_name: str
    product_type: str
    tech_requirements: list[str]
    process_rows: list[tuple[str, str, str]]


CASES = [
    Case(
        prefix="D125A-181200A003",
        part_name="Rotor Bracket",
        product_type="支架类",
        tech_requirements=["去毛刺", "关键尺寸全检", "表面无裂纹"],
        process_rows=[
            ("0010", "下料", "钢板下料，留加工余量"),
            ("0020", "粗加工", "铣平面并钻定位孔"),
            ("0030", "精加工", "精铣外形，保证尺寸要求"),
        ],
    ),
    Case(
        prefix="GJ-2024-001",
        part_name="Flange Support",
        product_type="法兰类",
        tech_requirements=["倒角去锐边", "重要孔位需复检"],
        process_rows=[
            ("0010", "Sawing", "Cut raw material to length"),
            ("0020", "Turning", "Rough turn outer diameter and end face"),
            ("0030", "Drilling", "Drill mounting holes to drawing"),
        ],
    ),
]


def _cleanup() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    XLSX_DIR.mkdir(parents=True, exist_ok=True)


def _font(size: int = 28):
    try:
        return ImageFont.truetype("arial.ttf", size=size)
    except Exception:
        return ImageFont.load_default()


def _make_pdf(case: Case) -> Path:
    img = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(img)
    title_font = _font(44)
    body_font = _font(28)
    small_font = _font(22)

    y = 90
    draw.text((70, y), f"DRAWING NO: {case.prefix}", fill="black", font=title_font)
    y += 90
    draw.text((70, y), f"PART NAME: {case.part_name}", fill="black", font=body_font)
    y += 60
    draw.text((70, y), f"PRODUCT TYPE: {case.product_type}", fill="black", font=body_font)
    y += 80
    draw.text((70, y), "TECHNICAL REQUIREMENTS:", fill="black", font=body_font)
    y += 50
    for item in case.tech_requirements:
        draw.text((110, y), f"- {item}", fill="black", font=small_font)
        y += 40

    y += 40
    draw.rectangle((60, y, 1180, y + 250), outline="black", width=3)
    draw.text((90, y + 25), "SAMPLE GEOMETRY / DIMENSION AREA", fill="black", font=body_font)
    draw.text((90, y + 85), "This synthetic page is used for ZIP import self-test.", fill="black", font=small_font)
    draw.text((90, y + 130), "It intentionally contains simple OCR-friendly labels.", fill="black", font=small_font)

    pdf_path = PDF_DIR / f"{case.prefix}.pdf"
    img.save(pdf_path, "PDF", resolution=144.0)
    return pdf_path


def _make_xlsx(case: Case) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "工艺表"
    ws.append(["图号", "工序内容"])
    for tag, name, content in case.process_rows:
        ws.append([case.prefix, f"{tag} {name}：{content}"])

    ws2 = wb.create_sheet("摘要")
    ws2.append(["图号", case.prefix])
    ws2.append(["零件名称", case.part_name])
    ws2.append(["产品类型", case.product_type])
    ws2.append(["技术要求", "；".join(case.tech_requirements)])
    ws2.append(["批次时间", datetime.now().isoformat()])

    xlsx_path = XLSX_DIR / f"{case.prefix}.xlsx"
    wb.save(xlsx_path)
    return xlsx_path


def _build_zip(files: list[Path]) -> None:
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file_path in files:
            arcname = file_path.relative_to(ROOT)
            zf.write(file_path, arcname.as_posix())


def main() -> None:
    _cleanup()

    generated: list[Path] = []
    manifest = {
        "created_at": datetime.now().isoformat(),
        "description": "Synthetic ZIP import self-test package for PDF + XLSX pairing",
        "cases": [],
    }

    for case in CASES:
        pdf_path = _make_pdf(case)
        xlsx_path = _make_xlsx(case)
        generated.extend([pdf_path, xlsx_path])
        manifest["cases"].append(
            {
                "prefix": case.prefix,
                "pdf": pdf_path.name,
                "xlsx": xlsx_path.name,
                "part_name": case.part_name,
            }
        )

    # Unmatched fixtures for negative-path testing
    extra_pdf = PDF_DIR / "UNMATCHED-AX99.pdf"
    extra_xlsx = XLSX_DIR / "UNMATCHED-BB99.xlsx"

    img = Image.new("RGB", (1240, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((80, 120), "UNMATCHED SAMPLE PDF", fill="black", font=_font(42))
    draw.text((80, 220), "This file is intentionally not paired with an XLSX.", fill="black", font=_font(26))
    img.save(extra_pdf, "PDF", resolution=144.0)

    wb = Workbook()
    ws = wb.active
    ws.title = "工艺表"
    ws.append(["图号", "UNMATCHED-BB99"])
    ws.append(["零件名称", "Unmatched Sample"])
    ws.append(["工序号", "工序名称", "工序内容"])
    ws.append(["0010", "检查", "This row is for unmatched-file testing only"])
    wb.save(extra_xlsx)

    generated.extend([extra_pdf, extra_xlsx])
    manifest["unmatched"] = [extra_pdf.name, extra_xlsx.name]

    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    generated.append(MANIFEST_PATH)
    _build_zip(generated)

    print(f"Generated: {ROOT}")
    print(f"ZIP: {ZIP_PATH}")
    print(f"Manifest: {MANIFEST_PATH}")


if __name__ == "__main__":
    main()
